# tests/test_paper.py

"""`docmarq.paper` around the markdown engine, and the shared parts."""

import dataclasses, inspect
from pathlib import Path
import docx, pytest
from docmarq import DOCX, RichSegment
from docmarq.constants import A4
from docmarq.parts import signatures, sign_labels, page_number, LANG_LABELS
from docmarq.paper import md_to_docx, md_into, PaperStyle, lang_style, LANG_PRESETS
from docmarq.md import MarkdownStyle, LANG_PRESETS as ENGINE_PRESETS, metadata_from_frontmatter
from conftest import NS, docx_xml, docx_text

# Style fields md2paper sets by name from request options (its `_STYLE` and `docx_style`).
MD2PAPER_FIELDS = {
  "body_size", "line_height", "image_max_h", "image_min_dpi", "banner", "banner_compact",
  "sign", "mermaid_theme", "link_root", "link_base", "font_body", "font_head", "font_mono",
  "page_number_label",
}

BANNER = "---\ntitle: Report\nid: R-1\nauthor: Ann\n---\n\n# Report\n\nBody text.\n"

#-------------------------------------------------------------------------------- md2paper contract

def style_carries_every_md2paper_field():
  fields = {f.name for f in dataclasses.fields(PaperStyle)}
  assert MD2PAPER_FIELDS <= fields, MD2PAPER_FIELDS - fields

def md_to_docx_signature_is_stable():
  params = inspect.signature(md_to_docx).parameters
  assert list(params) == ["md_text", "output_path", "style", "page", "margin", "gutter",
    "base_dir", "font_dir", "metadata"]
  assert (params["page"].default, params["margin"].default, params["gutter"].default) \
    == (A4, 20, 0)

def lang_style_layers_engine_parts_and_banner_labels():
  style = lang_style("de")
  assert isinstance(style, PaperStyle)
  assert (style.page_number_label, style.callout_label_note, style.banner_label_created) \
    == ("Seite", "Hinweis", "Erstellt")
  for lang, preset in LANG_PRESETS.items():
    assert ENGINE_PRESETS[lang].items() <= preset.items()
    assert LANG_LABELS.get(lang, {}).items() <= preset.items()

def lang_style_copies_its_preset():
  lang_style("pl").sign_labels["signature"].append("changed")
  assert lang_style("pl").sign_labels["signature"] == ["Podpis"]

def engine_style_is_refused(tmp_path):
  with pytest.raises(TypeError, match="PaperStyle"):
    md_to_docx("# Title", str(tmp_path / "x.docx"), style=MarkdownStyle())

#----------------------------------------------------------------------------------------- Document

def banner_takes_the_title_and_the_compact_header_the_id(tmp_path):
  path = str(tmp_path / "doc.docx")
  md_to_docx(BANNER, path)
  text = docx_text(path)
  assert text.count("Report") == 1 # the repeated `# Report` is dropped
  headers = [docx_xml(path, f"word/header{i}.xml") for i in (1, 2)]
  assert ["R-1" in "".join(h.xpath("//w:t/text()", namespaces=NS)) for h in headers] \
    == [False, True]

def page_number_sits_in_both_footers_under_a_banner(tmp_path):
  path = str(tmp_path / "doc.docx")
  md_to_docx(BANNER, path)
  for i in (1, 2):
    footer = docx_xml(path, f"word/footer{i}.xml")
    assert footer.xpath("//w:fldSimple/@w:instr", namespaces=NS) == ["PAGE", "NUMPAGES"]

def no_page_number_without_a_label(tmp_path):
  path = str(tmp_path / "doc.docx")
  md_to_docx("Body.", path, style=PaperStyle(page_number_label=None))
  import zipfile
  assert not [n for n in zipfile.ZipFile(path).namelist() if "footer" in n]

def signatures_end_the_document(tmp_path):
  path = str(tmp_path / "doc.docx")
  md_to_docx("Body.", path, style=lang_style("pl", sign="contract"))
  root = docx_xml(path)
  assert root.xpath("//w:tbl//w:t/text()", namespaces=NS) == ["Zleceniodawca", "Zleceniobiorca"]
  assert root.xpath("//w:tbl//w:pBdr/w:bottom/@w:val", namespaces=NS) == ["single", "single"]

#-------------------------------------------------------------------------------------------- Parts

def sign_labels_read_scenarios_and_lists():
  scenarios = {"signature": ["Signature"], "approval": ["Prepared", "Approved"]}
  assert sign_labels(True, scenarios) == ["Signature"]
  assert sign_labels("approval", scenarios) == ["Prepared", "Approved"]
  assert sign_labels(["A", "B"], scenarios) == ["A", "B"]
  assert sign_labels("custom", scenarios) == ["custom"]
  assert sign_labels(True, {}) == [""]
  assert sign_labels(False, scenarios) == []

def one_signing_line_sits_at_the_right_margin(tmp_path):
  doc = DOCX(str(tmp_path / "s.docx"))
  face = RichSegment("", size=9, color=(0.2, 0.2, 0.2), italic=True)
  signatures(doc, ["Name"], face, gap=20, width=70, thick=0.5, line="dotted")
  doc.save()
  rule = docx_xml(doc.path).xpath("//w:p[w:pPr/w:pBdr]", namespaces=NS)[0]
  assert rule.xpath("string(w:pPr/w:pBdr/w:bottom/@w:val)", namespaces=NS) == "dotted"
  assert int(rule.xpath("string(w:pPr/w:ind/@w:left)", namespaces=NS)) \
    == round((170 - 70) * 1440 / 25.4)

def page_number_takes_a_face(tmp_path):
  doc = DOCX(str(tmp_path / "n.docx"))
  p = doc.doc.sections[0].footer.paragraphs[0]
  page_number(p, "Strona", show_total=False, face=RichSegment("", family="Cambria", size=8))
  doc.save()
  footer = docx_xml(doc.path, "word/footer1.xml")
  assert footer.xpath("//w:fldSimple/@w:instr", namespaces=NS) == ["PAGE"]
  assert set(footer.xpath("//w:r/w:rPr/w:sz/@w:val", namespaces=NS)) == {"16"}

def page_number_reads_page_and_total(tmp_path):
  path = str(tmp_path / "doc.docx")
  md_to_docx("Body.", path)
  footer = docx_xml(path, "word/footer1.xml")
  assert footer.xpath("string(//w:p)", namespaces=NS) == "Page 0/0" # fields show 0 until Word
  size = str(round(PaperStyle().banner_meta_size * 2)) # half-points
  assert set(footer.xpath("//w:r/w:rPr/w:sz/@w:val", namespaces=NS)) == {size}

#------------------------------------------------------------------------------------------ md_into

def md_into_writes_after_a_cover(tmp_path):
  doc = DOCX(str(tmp_path / "cover.docx"))
  doc.para("Coverword").page_break()
  md_into(doc, BANNER, style=PaperStyle())
  doc.save()
  texts = docx_xml(doc.path).xpath("//w:t/text()", namespaces=NS)
  assert texts[0] == "Coverword" and "Body text." in texts
  assert docx.Document(doc.path).core_properties.title == "Report"
  footer = docx_xml(doc.path, "word/footer1.xml")
  assert footer.xpath("//w:fldSimple/@w:instr", namespaces=NS) == ["PAGE", "NUMPAGES"]

def missing_logo_warns_at_the_callers_line(tmp_path):
  md = "---\ntitle: T\nlogo: none.png\n---\n\nBody.\n"
  with pytest.warns(RuntimeWarning, match="not found") as caught:
    md_to_docx(md, str(tmp_path / "a.docx"))
    md_into(DOCX(str(tmp_path / "b.docx")), md)
  assert {Path(w.filename).name for w in caught} == {Path(__file__).name}

def invalid_frontmatter_warns_once(tmp_path):
  with pytest.warns(RuntimeWarning, match="invalid YAML") as caught:
    md_to_docx("---\ntitle: [\n---\n\nBody.\n", str(tmp_path / "y.docx"))
  assert len([w for w in caught if "invalid YAML" in str(w.message)]) == 1

def title_given_as_a_list_is_one_line(tmp_path):
  fm = {"title": ["Umowa", "o dzieło"], "author": "Jan\nKowalski"}
  assert metadata_from_frontmatter(fm) == {"title": "Umowa o dzieło", "author": "Jan Kowalski"}
  path = str(tmp_path / "t.docx")
  md_to_docx("---\ntitle: [Umowa, o dzieło]\n---\n\nBody.\n", path)
  assert docx.Document(path).core_properties.title == "Umowa o dzieło"

def status_given_as_a_number_is_a_badge(tmp_path):
  path = str(tmp_path / "s.docx")
  md_to_docx("---\ntitle: T\nstatus: 1\n---\n\nBody.\n", path)
  assert docx_xml(path).xpath("//w:t[.=' 1 ']", namespaces=NS)
