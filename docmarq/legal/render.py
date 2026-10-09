# docmarq/legal/render.py

"""
Letters, statements and contracts: the page around the markdown content.

`md_to_docx` lays out one document, `md_parallel_to_docx` versions of one document side by side.
Word lays the pages out: it breaks them, balances the columns and numbers the pages.
"""
from dataclasses import dataclass, replace
from markdown_it.token import Token
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..core import DOCX, _set_table_widths, _mark_language, _no_hyphens
from ..constants import PageSize, A4
from ..inline import RichSegment
from ..parts import signatures, page_number
from ..tables import remove_table_borders, set_cell_margins
from ..utils import mm_to_pt
from ..md import (MarkdownRenderer, split_frontmatter, docx_metadata, top_blocks,
  is_pagebreak_directive, is_group_open_directive)
from .style import LegalStyle
from .header import draw_header, draw_attachments, lines_of, has_header

#--------------------------------------------------------------------------------------------- Page

def md_to_docx(
  md_text:str,
  output_path:str,
  *,
  style:LegalStyle|None = None,
  page:PageSize = A4,
  margin:float|tuple = 20,
  gutter:float = 0,
  base_dir:str|None = None,
  font_dir:str|None = None,
  metadata:dict|None = None,
) -> DOCX:
  """
  Letter, statement or contract: markdown content with a letter header and a signed end.

  The text before the first `##` is the preamble, the rest the clauses.
  The arguments match `docmarq.paper.md_to_docx`.
  """
  style = _checked(style or LegalStyle())
  fm, _ = split_frontmatter(md_text)
  doc = _new_docx(output_path, page, margin, gutter, style, fm, metadata)
  renderer = MarkdownRenderer(doc, style, base_dir=base_dir, font_dir=font_dir)
  _footer(doc, style)
  floats = draw_header(doc, style, fm, renderer.frontmatter_file(fm, "logo"))
  tokens = renderer.parse(md_text)
  renderer.index(tokens)
  first = _clauses(tokens)
  _preamble(doc, renderer, style, tokens[:first], floats)
  # a continuous break ends the columns, so Word balances them, on the last page too
  doc.columns(style.columns, gap=style.column_gap)
  renderer.render_tokens(tokens[first:])
  doc.columns(1)
  labels, items = style.sign_lines(), lines_of(fm.get("attachments"))
  if labels or items:
    _tail(doc, style, lambda: _end(doc, style, labels, items))
  doc.save()
  return doc

def _checked(style) -> LegalStyle:
  """`style` when it is a usable `LegalStyle`. A wrong one fails here, before a word is written."""
  if not isinstance(style, LegalStyle):
    raise TypeError("style must be a LegalStyle, build one with docmarq.legal.lang_style()")
  if style.logo_align not in ("L", "C", "R"):
    raise ValueError(f"logo_align must be L, C or R, got {style.logo_align!r}")
  if style.to_align not in ("L", "R"):
    raise ValueError(f"to_align must be L or R, got {style.to_align!r}")
  if not isinstance(style.columns, int) or style.columns < 1:
    raise ValueError(f"columns must be a whole number from 1 up, got {style.columns!r}")
  return style

def _new_docx(
  output_path:str, page:PageSize, margin:float|tuple, gutter:float,
  style:LegalStyle, fm:dict, metadata:dict|None,
) -> DOCX:
  """The output `DOCX`, core properties from frontmatter, which `metadata=` overrides per key."""
  doc = DOCX(output_path, width=page.width, height=page.height,
    margin=margin, gutter=gutter,
    font_body=style.font_body, font_head=style.font_head, body_size=style.body_size)
  # two lines at least on each side of a page break
  doc.doc.styles["Normal"].paragraph_format.widow_control = True
  # text in the style's colour, a paragraph spaced by the gap below the one before
  doc.style(space_before=0, color=style.body_color)
  meta = docx_metadata(fm, metadata)
  if meta: doc.metadata(**meta)
  return doc

def _footer(doc:DOCX, style:LegalStyle):
  """Page number on every page of a document longer than one."""
  if not style.page_number_label: return
  face = RichSegment("", family=style.page_font, size=style.header_size, color=style.muted_color)
  page_number(doc.doc.sections[-1].footer.paragraphs[0], style.page_number_label,
    show_total=style.page_number_total, face=face, many=True)

#----------------------------------------------------------------------------------------- Preamble

def _clauses(tokens:list[Token]) -> int:
  """Start of the first top-level `##`, where the clauses begin. Without one it is `0`."""
  return next((start for start, _ in top_blocks(tokens)
    if tokens[start].type == "heading_open" and tokens[start].tag == "h2"), 0)

def _preamble(
  doc:DOCX, renderer:MarkdownRenderer, style:LegalStyle,
  tokens:list[Token], floats:bool,
):
  """
  The preamble in the header face and its own size and spacing,
  beside the logo when it floats by the title.
  The clauses start `preamble_gap` below the preamble and the logo.
  """
  if not tokens and not floats: return
  look = replace(style, font_body=style.page_font,
    body_size=style.preamble_size or style.body_size,
    para_gap=style.para_gap if style.preamble_para_gap is None else style.preamble_para_gap)
  renderer.render_tokens(tokens, look)
  if floats: doc._hairline().add_run().add_break(WD_BREAK.LINE_CLEAR_ALL)
  # the clauses start as after a paragraph of the preamble, `preamble_gap` lower
  doc._track_block_spacing(mm_to_pt(look.para_gap), extra_for_next=mm_to_pt(style.preamble_gap))

#---------------------------------------------------------------------------------------------- End

def _end(doc:DOCX, style:LegalStyle, labels:list[str], items:list[str]):
  """Dotted signing lines in the text colour, labels in the header face, then the attachments."""
  face = RichSegment("", family=style.page_font, size=style.sign_size, color=style.body_color,
    italic=True)
  signatures(doc, labels, face, gap=style.sign_gap, width=style.sign_w, thick=style.hr_thick,
    line="dotted")
  draw_attachments(doc, style, items)

def _tail(doc:DOCX, style:LegalStyle, write):
  """
  `write()` the end of the document, kept together on one page.
  With `sign_bottom` it floats at the bottom margin of the page the text ends on.
  """
  if style.sign_bottom:
    cell, width = _bottom_cell(doc)
    with doc.into(cell, width):
      write()
    doc._hairline()
    return
  body = doc.doc.element.body
  start = len(body) - 1 # the body ends with its `w:sectPr`
  write()
  paras = [p for el in list(body)[start:] for p in el.iter(qn("w:p"))]
  for p in paras[:-1]:
    p.get_or_add_pPr().keepNext_val = True

def _bottom_cell(doc:DOCX) -> tuple:
  """`(cell, width)` of a borderless table one text width wide, floating at the bottom margin."""
  width = doc.content_width
  doc._flush_para()
  tbl = doc._new_table(1, 1)
  _set_table_widths(tbl, [width])
  remove_table_borders(tbl)
  cell = tbl.rows[0].cells[0]
  set_cell_margins(cell, top=0, right=0, bot=0, left=0)
  place = OxmlElement("w:tblpPr", {qn("w:vertAnchor"): "margin", qn("w:horzAnchor"): "margin",
    qn("w:tblpXSpec"): "center", qn("w:tblpYSpec"): "bottom"})
  tbl._tbl.tblPr.insert(0, place)
  place.addnext(OxmlElement("w:tblOverlap", {qn("w:val"): "never"}))
  return cell, width

#----------------------------------------------------------------------------------------- Parallel

def md_parallel_to_docx(
  md_texts:list[str],
  output_path:str,
  *,
  styles:list[LegalStyle],
  page:PageSize = A4,
  margin:float|tuple = 20,
  gutter:float = 0,
  base_dir:str|None = None,
  font_dir:str|None = None,
  metadata:dict|None = None,
) -> DOCX:
  """
  Versions of one document side by side, one column each, usually one per language.

  Each version is an ordinary legal document written with its own style,
  so its own labels, date format and language.
  Which language a version is in is the caller's knowledge, not this one's.
  A borderless table holds them: a column per version, a row per top-level block,
  so every row starts level in all columns.
  Word may break a row taller than the rest of a page.
  Signing lines, the page number and the heading styles come from `styles[0]`,
  the core properties from the first version's frontmatter, which `metadata=` overrides per key.
  """
  if len(md_texts) < 2:
    raise ValueError(f"a parallel document needs two versions or more, got {len(md_texts)}")
  if len(styles) != len(md_texts):
    raise ValueError(f"{len(md_texts)} versions need as many styles, got {len(styles)}")
  styles = [_checked(s) for s in styles]
  fronts = [split_frontmatter(text)[0] for text in md_texts]
  doc = _new_docx(output_path, page, margin, gutter, styles[0], fronts[0], metadata)
  # built last to first: each sets the document up in its own style, and the first one's holds
  versions = [_version(doc, styles[i], md_texts[i], fronts[i], base_dir, font_dir, i)
    for i in reversed(range(len(styles)))][::-1]
  rows = _pair([v.tokens for v in versions])
  _footer(doc, styles[0])
  grid = _Grid(doc, styles)
  if any(has_header(v.front, v.logo) for v in versions):
    grid.row(lambda i: draw_header(doc, styles[i], fronts[i], versions[i].logo, beside=False))
  for k, row in enumerate(rows):
    kind = _kind(versions[0].tokens, row[0][0])
    if kind == "pagebreak":
      grid.page_break()
      continue
    tr = grid.row(lambda i: versions[i].renderer.render_tokens(
      versions[i].tokens[slice(*row[i])], styles[i]))
    # a heading row keeps the row after it company, unless a page break comes between
    follows = _kind(versions[0].tokens, rows[k + 1][0][0]) if k + 1 < len(rows) else "pagebreak"
    if kind.startswith("heading_open") and follows != "pagebreak":
      for p in tr.iter(qn("w:p")):
        p.get_or_add_pPr().keepNext_val = True
  items = [lines_of(fm.get("attachments")) for fm in fronts]
  labels = styles[0].sign_lines()
  if labels or any(items):
    _tail(doc, styles[0], lambda: _parallel_end(doc, styles, labels, items))
  doc.save()
  return doc

@dataclass
class _Version:
  """One version of a parallel document: its facts, its text and what writes it."""
  front: dict
  logo: str|None
  renderer: MarkdownRenderer
  tokens: list[Token]

def _version(
  doc:DOCX, style:LegalStyle, md_text:str, fm:dict,
  base_dir:str|None, font_dir:str|None, i:int,
) -> _Version:
  """
  Version `i` parsed and indexed, its bookmarks named apart from the other versions' bookmarks,
  so a link in one version jumps to a heading of its own.
  """
  renderer = MarkdownRenderer(doc, style, base_dir=base_dir, font_dir=font_dir)
  tokens = renderer.parse(md_text)
  renderer.index(tokens, prefix=f"v{i + 1}-" if i else "")
  return _Version(fm, renderer.frontmatter_file(fm, "logo"), renderer, tokens)

class _Grid:
  """
  Borderless table of the versions, a column each `column_gap` apart, rows added one by one.
  A page break ends it and starts the next on a new page.
  """
  def __init__(self, doc:DOCX, styles:list[LegalStyle]):
    self.doc = doc
    self.styles = styles
    self.count = len(styles)
    self.gap = styles[0].column_gap
    self.width = (doc.content_width - self.gap * (self.count - 1)) / self.count
    self.table = None

  def row(self, write):
    """New row, `write(i)` called in the cell of each version `i`. Returns the row's `w:tr`."""
    if self.table is None:
      self.table = _grid_table(self.doc, self.count, self.width, self.gap)
      cells = self.table.rows[0].cells
    else:
      cells = self.table.add_row().cells
    half = self.gap / 2
    for i, cell in enumerate(cells):
      set_cell_margins(cell, top=0, bot=0,
        left=half if i else 0, right=half if i < self.count - 1 else 0)
      with self.doc.into(cell, self.width):
        # a row follows the one above as a paragraph does: its first gap collapses into theirs
        self.doc._track_block_spacing(mm_to_pt(self.styles[i].para_gap))
        write(i)
      lang = self.styles[i].hyphenate
      if lang: _mark_language(cell._tc, lang)
      # a version naming no language is not split by the rules of another's
      elif self.doc._hyphenate:
        for p in cell._tc.iter(qn("w:p")): _no_hyphens(p)
    return cells[0]._tc.getparent()

  def page_break(self):
    """End the table, so the rows after a page break start a new one on the next page."""
    if self.table is None: return
    self.table = None
    self.doc._hairline().paragraph_format.page_break_before = True

def _grid_table(doc:DOCX, count:int, width:float, gap:float):
  """Borderless one-row table of `count` columns `width` mm wide, `gap` mm apart."""
  doc._flush_para()
  tbl = doc._new_table(1, count)
  half = gap / 2
  _set_table_widths(tbl, [width + (half if i else 0) + (half if i < count - 1 else 0)
    for i in range(count)])
  remove_table_borders(tbl)
  return tbl

def _kind(tokens:list[Token], start:int) -> str:
  """What a top-level block is, for matching versions: its opening token, a directive by name."""
  t = tokens[start]
  if t.type == "html_block":
    if is_pagebreak_directive(t.content or ""): return "pagebreak"
    if is_group_open_directive(t.content or ""): return "group"
  return f"{t.type} {t.tag}".strip()

def _preview(tokens:list[Token], span:tuple[int, int]|None) -> str:
  """Start of the text of the block at `span`, for an error message."""
  if span is None: return ""
  start, end = span
  text = next((t.content for t in tokens[start:end] if t.type == "inline"), "")
  text = " ".join(text.split())
  return f" {text[:40] + '...' if len(text) > 40 else text!r}" if text else ""

def _pair(versions:list[list[Token]]) -> list[list[tuple[int, int]]]:
  """
  Rows of block spans, one span per version, after checking the versions line up.
  Differing block counts or kinds raise, naming the first block where they part ways,
  and so does a page break inside a block, which would break a page mid-row.
  """
  blocks = [top_blocks(tokens) for tokens in versions]
  for k in range(max(len(b) for b in blocks)):
    spans = [b[k] if k < len(b) else None for b in blocks]
    kinds = [_kind(t, s[0]) if s else "nothing" for t, s in zip(versions, spans)]
    if len(set(kinds)) > 1:
      parts = "; ".join(f"version {i}: {kind}{_preview(t, s)}"
        for i, (kind, t, s) in enumerate(zip(kinds, versions, spans), 1))
      raise ValueError(f"versions part ways at block {k + 1}: {parts}")
  for tokens, spans in zip(versions, blocks):
    for start, end in spans:
      if any(t.type == "html_block" and is_pagebreak_directive(t.content or "")
          for t in tokens[start + 1:end]):
        raise ValueError("a <!-- pagebreak --> inside a block would break a page mid-row")
  return [list(row) for row in zip(*blocks)]

def _parallel_end(doc:DOCX, styles:list[LegalStyle], labels:list[str], items:list[list[str]]):
  """Signing lines of `styles[0]` across the page, then the attachments of each version below."""
  first = styles[0]
  face = RichSegment("", family=first.page_font, size=first.sign_size, color=first.body_color,
    italic=True)
  signatures(doc, labels, face, gap=first.sign_gap, width=first.sign_w, thick=first.hr_thick,
    line="dotted")
  if any(items):
    _Grid(doc, styles).row(lambda i: draw_attachments(doc, styles[i], items[i]))
