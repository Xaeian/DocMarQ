# docmarq/legal/header.py

"""
Letter header and attachment list of a legal document, both written from frontmatter.

Header rows, each written only when its keys are set:

  1. `logo`, aligned by `style.logo_align`
  2. `from` on the left, `place` and `date` on the right of its first line
  3. `to` flush with the right margin, or the left one with `style.to_align`,
     then `ref` under it: the case number, in bold after `style.ref_label`
  4. `title`, centred

A logo at the side shares the title's row when nothing stands between them, as on a contract:
it floats at the margin, and the preamble after the title runs beside it.
Values are plain text. A YAML block (`|`) or a list gives one line per entry,
or entries joined by commas for `place` and `ref`, which keep to one line.
"""
import io, warnings
from docx.shared import Mm, Pt
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from ..core import DOCX
from ..constants import EMU_PER_MM
from ..metrics import text_width_mm
from ..svg import svg_to_png_buffer, is_svg
from ..tables import remove_table_borders, set_cell_margins, set_table_widths
from ..utils import format_date, mm_to_pt, align_to_docx
from .._warn import caller_level
from .style import LegalStyle

_ROW_GAP_MM = 6 # below every header row
_LIST_GAP_MM = 6 # above the attachments
_SIDE_GAP_MM = 5 # between text and what stands beside it: the date, a logo
_TITLE_LEADING = 1.2 # title line height, of its size

#------------------------------------------------------------------------------------------- Values

def lines_of(value) -> list[str]:
  """
  A YAML block text or list as lines, blank and empty entries dropped. `None` gives none.
  A mapping is refused: no header value has named parts.
  """
  if value is None: return []
  if isinstance(value, dict):
    raise ValueError(f"a frontmatter value here is text or a list, got a mapping: {value!r}")
  items = value if isinstance(value, (list, tuple)) else str(value).splitlines()
  return [str(item).strip() for item in items if item is not None and str(item).strip()]

def _one_line(value) -> str:
  """A YAML value on one line: the entries of a block or a list joined by commas."""
  return ", ".join(lines_of(value))

def _face(style:LegalStyle, bold:bool=False) -> dict:
  """`DOCX.text()` arguments for header-size text in the face of the page around the text."""
  return {"family": style.page_font, "size": style.header_size, "color": style.body_color,
    "bold": bold}

def _width(style:LegalStyle, text:str) -> float:
  """Width (mm) of `text` in the header face."""
  return text_width_mm(text, style.page_font, style.header_size)

def _lines(
  doc:DOCX, style:LegalStyle, lines:list,
  align:str, after:float,
  indent:tuple[float, float] = (0, 0),
):
  """
  One paragraph after the last block, holding `lines` one under another at header size.
  A line is a text, or `(text, bold)` pairs. `after` is the space below in mm,
  `indent` the left and right indents in mm.
  """
  doc.para(align=align)
  p = doc.paragraph
  pf = p.paragraph_format
  pf.space_before = Pt(0)
  pf.space_after = Pt(mm_to_pt(after))
  if indent[0]: pf.left_indent = Mm(indent[0])
  if indent[1]: pf.right_indent = Mm(indent[1])
  doc.keep_words(p)
  for i, line in enumerate(lines):
    if i: doc.line_break()
    for text, bold in [(line, False)] if isinstance(line, str) else line:
      doc.text(text, **_face(style, bold))
  doc.track_block_spacing(mm_to_pt(after))

#------------------------------------------------------------------------------------------- Header

def has_header(fm:dict, logo:str|None) -> bool:
  """True when `draw_header` has anything to write: a logo or a frontmatter fact of the header."""
  return bool(logo) or any(fm.get(key) for key in ("title", "place", "date", "from", "to", "ref"))

def draw_header(
  doc:DOCX, style:LegalStyle, fm:dict, logo:str|None,
  beside:bool = True,
) -> bool:
  """
  Write the header rows after the last block.

  `logo` is the resolved path of an existing file, or `None`.
  Returns `True` when the logo floats beside a lone title:
  the preamble then runs beside it, and the caller ends it below the logo.
  `beside=False` keeps the logo in a row of its own, as in a column of a parallel document.
  """
  width = doc.content_width
  sender = lines_of(fm.get("from"))
  place = _one_line(fm.get("place"))
  stamp = ", ".join(t for t in (place, format_date(fm.get("date"), style.date_format)) if t)
  recipient: list = lines_of(fm.get("to"))
  ref = _one_line(fm.get("ref"))
  if ref:
    recipient.append([(f"{style.ref_label}: ", False), (ref, True)])
  title = lines_of(fm.get("title"))
  pic = _picture(style, logo) if logo else None
  # beside a lone title only, kept clear on both sides so the title stays centred
  clear = pic[1] + _SIDE_GAP_MM if pic else 0
  floats = bool(beside and pic and title and style.logo_align != "C"
    and not (sender or stamp or recipient)
    and width - 2 * clear >= _widest_word(style, title))
  rows = False
  if pic and not floats:
    _logo_row(doc, style, pic)
    rows = True
  if sender or stamp:
    _sender_row(doc, style, sender, stamp, width)
    rows = True
  if recipient:
    half = (width / 2, 0) if style.to_align == "R" else (0, width / 2)
    _lines(doc, style, recipient, style.to_align, _ROW_GAP_MM, indent=half)
    rows = True
  if title:
    _title(doc, style, title, rows, pic if floats else None, clear)
  return floats

def _sender_row(doc:DOCX, style:LegalStyle, sender:list[str], stamp:str, width:float):
  """
  The sender on the left, place and date on the right of its first line.
  Beside the stamp the sender wraps in the room left; with too little room for a word,
  the stamp takes a line of its own above the sender.
  """
  if not sender or not stamp:
    _lines(doc, style, sender or [stamp], "L" if sender else "R", _ROW_GAP_MM)
    return
  right = _width(style, stamp) + _SIDE_GAP_MM
  if width - right < max(_width(style, w) for line in sender for w in line.split()):
    _lines(doc, style, [stamp], "R", 0)
    _lines(doc, style, sender, "L", _ROW_GAP_MM)
    return
  doc.enter()
  tbl = doc.new_table(1, 2)
  set_table_widths(tbl, [width - right, right])
  remove_table_borders(tbl)
  left, side = tbl.rows[0].cells
  for cell in (left, side):
    set_cell_margins(cell, top=0, right=0, bot=0, left=0)
  with doc.into(left, width - right):
    _lines(doc, style, sender, "L", _ROW_GAP_MM)
  with doc.into(side, right):
    _lines(doc, style, [stamp], "R", _ROW_GAP_MM)
  doc.track_block_spacing(0)

def _title(
  doc:DOCX, style:LegalStyle, lines:list[str], rows:bool,
  pic:tuple|None, clear:float,
):
  """
  The title centred, a line per entry, in the header face or the heading face when none is set.
  With `title_caps` its lowercase letters are Word's small capitals.
  With `pic` the logo floats beside it, the title indented `clear` mm on both sides.
  """
  doc.para(align="C")
  p = doc.paragraph
  pf = p.paragraph_format
  # under header rows the gap above is `title_gap`, the row gap included
  pf.space_before = Pt(mm_to_pt(max(0, style.title_gap - _ROW_GAP_MM)) if rows else 0)
  pf.space_after = Pt(mm_to_pt(style.title_gap))
  pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
  pf.line_spacing = Pt(style.title_size * _TITLE_LEADING)
  doc.keep_words(p)
  if pic:
    pf.left_indent = pf.right_indent = Mm(clear)
    _float(p.add_run(), pic, style.logo_align)
  mode = style.head_mode.lower()
  face = {"family": style.header_font or style.font_head, "size": style.title_size,
    "color": style.head_color, "bold": "bold" in mode, "italic": "italic" in mode,
    "small_caps": style.title_caps}
  for i, line in enumerate(lines):
    if i: doc.line_break()
    doc.text(line, **face)
  doc.track_block_spacing(mm_to_pt(style.title_gap))

def _widest_word(style:LegalStyle, lines:list[str]) -> float:
  """Width (mm) of the widest word of the title, capitals standing in for small capitals."""
  family = style.header_font or style.font_head
  bold = "bold" in style.head_mode.lower()
  words = [w.upper() if style.title_caps else w for line in lines for w in line.split()]
  return max((text_width_mm(w, family, style.title_size, bold) for w in words), default=0)

#--------------------------------------------------------------------------------------------- Logo

def _picture(style:LegalStyle, path:str) -> tuple|None:
  """
  `(source, width, height)` of the logo in mm, within `logo_max_w` × `logo_max_h`, aspect kept.
  It goes to Word as a PNG, whatever it was: Word takes no WebP or CMYK.
  `None`, with a warning, for a file that will not open.
  """
  try:
    src = svg_to_png_buffer(path) if is_svg(path) else path
    if src is None: raise OSError("no SVG backend")
    with Image.open(src) as im:
      aspect = im.width / im.height
      png = io.BytesIO()
      im.convert("RGBA" if "A" in im.getbands() or im.mode == "P" else "RGB").save(png, "PNG")
  except OSError as e:
    warnings.warn(f"frontmatter logo {path!r} could not be read, written without it: {e}",
      RuntimeWarning, stacklevel=caller_level())
    return None
  src = png
  src.seek(0)
  w, h = style.logo_max_h * aspect, style.logo_max_h
  if w > style.logo_max_w:
    w, h = style.logo_max_w, style.logo_max_w / aspect
  return src, w, h

def _logo_row(doc:DOCX, style:LegalStyle, pic:tuple):
  """Logo in a paragraph of its own, aligned by `logo_align`, the row gap below."""
  src, w, h = pic
  doc.enter()
  p = doc.new_para()
  p.alignment = align_to_docx(style.logo_align)
  pf = p.paragraph_format
  pf.space_before = Pt(0)
  pf.space_after = Pt(mm_to_pt(_ROW_GAP_MM))
  p.add_run().add_picture(src, width=Mm(w), height=Mm(h))
  doc.track_block_spacing(mm_to_pt(_ROW_GAP_MM))

def _float(run, pic:tuple, align:str):
  """
  The logo in `run`, floating at the `align` side of the margin from the top of its paragraph.
  Text wraps square around it, `_SIDE_GAP_MM` clear on the side facing the text.
  """
  src, w, h = pic
  run.add_picture(src, width=Mm(w), height=Mm(h))
  inline = run._r.find(f".//{qn('wp:inline')}")
  gap = str(round(_SIDE_GAP_MM * EMU_PER_MM))
  anchor = OxmlElement("wp:anchor", {"distT": "0", "distB": "0",
    "distL": gap if align == "R" else "0", "distR": gap if align == "L" else "0",
    "simplePos": "0", "relativeHeight": "1", "behindDoc": "0", "locked": "0",
    "layoutInCell": "1", "allowOverlap": "0"})
  anchor.append(OxmlElement("wp:simplePos", {"x": "0", "y": "0"}))
  horizontal = OxmlElement("wp:positionH", {"relativeFrom": "margin"})
  side = OxmlElement("wp:align")
  side.text = "left" if align == "L" else "right"
  horizontal.append(side)
  anchor.append(horizontal)
  vertical = OxmlElement("wp:positionV", {"relativeFrom": "paragraph"})
  offset = OxmlElement("wp:posOffset")
  offset.text = "0"
  vertical.append(offset)
  anchor.append(vertical)
  anchor.append(inline.find(qn("wp:extent")))
  anchor.append(OxmlElement("wp:effectExtent", {"l": "0", "t": "0", "r": "0", "b": "0"}))
  anchor.append(OxmlElement("wp:wrapSquare", {"wrapText": "bothSides"}))
  for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
    anchor.append(inline.find(qn(tag)))
  inline.getparent().replace(inline, anchor)

#-------------------------------------------------------------------------------------- Attachments

def draw_attachments(doc:DOCX, style:LegalStyle, items:list[str]):
  """Attachment list after the last block: its label in bold, then the items numbered."""
  if not items: return
  doc.para()
  p = doc.paragraph
  pf = p.paragraph_format
  pf.space_before = Pt(mm_to_pt(_LIST_GAP_MM))
  pf.space_after = Pt(0)
  doc.keep_words(p)
  doc.text(f"{style.attachments_label}:", **_face(style, bold=True))
  for i, item in enumerate(items, 1):
    doc.line_break()
    doc.text(f"{i}. {item}", **_face(style))
  doc.track_block_spacing(0)
