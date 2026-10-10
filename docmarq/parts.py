# docmarq/parts.py

"""
Document parts built with the core API: signing lines and the page number,
with the labels and date format they take in each language.

They need no markdown, so a caller passes its look in as plain values.
Lengths are mm: the `DOCX` they write into must use mm, as the markdown engine does.
"""
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .inline import RichSegment, _apply_run_format
from .utils import mm_to_pt, color_hex, eighths
from .fields import page_fields, page_fields_if_many
from .tables import remove_table_borders, set_cell_margins, set_table_widths
from .core import _set_pbdr

_SIGN_SPACING_MM = 6 # least space between two signing lines

#------------------------------------------------------------------------------------------- Labels

def default_sign_labels() -> dict:
  """Signing scenarios in English: name → one label per signing line."""
  return {
    "signature": ["Signature"],
    "approval": ["Prepared by", "Approved by"],
    "contract": ["Client", "Contractor"],
  }

# Page number label, date format and signing scenarios per language, English being the default.
# The presets of the document packages build on them.
LANG_LABELS: dict[str, dict] = {
  "pl": {
    "page_number_label": "Strona",
    "date_format": "%d.%m.%Y",
    "sign_labels": {"signature": ["Podpis"], "approval": ["Opracował", "Zatwierdził"],
      "contract": ["Zleceniodawca", "Zleceniobiorca"]},
  },
  "de": {
    "page_number_label": "Seite",
    "date_format": "%d.%m.%Y",
    "sign_labels": {"signature": ["Unterschrift"], "approval": ["Erstellt", "Genehmigt"],
      "contract": ["Auftraggeber", "Auftragnehmer"]},
  },
  "fr": {
    "page_number_label": "Page",
    "date_format": "%d/%m/%Y",
    "sign_labels": {"signature": ["Signature"], "approval": ["Rédigé par", "Approuvé par"],
      "contract": ["Donneur d'ordre", "Prestataire"]},
  },
  "es": {
    "page_number_label": "Página",
    "date_format": "%d/%m/%Y",
    "sign_labels": {"signature": ["Firma"], "approval": ["Elaborado por", "Aprobado por"],
      "contract": ["Cliente", "Contratista"]},
  },
  "it": {
    "page_number_label": "Pagina",
    "date_format": "%d/%m/%Y",
    "sign_labels": {"signature": ["Firma"], "approval": ["Redatto da", "Approvato da"],
      "contract": ["Committente", "Appaltatore"]},
  },
  "cs": {
    "page_number_label": "Strana",
    "date_format": "%d.%m.%Y",
    "sign_labels": {"signature": ["Podpis"], "approval": ["Vypracoval", "Schválil"],
      "contract": ["Objednatel", "Zhotovitel"]},
  },
  "sk": {
    "page_number_label": "Strana",
    "date_format": "%d.%m.%Y",
    "sign_labels": {"signature": ["Podpis"], "approval": ["Vypracoval", "Schválil"],
      "contract": ["Objednávateľ", "Zhotoviteľ"]},
  },
}

def sign_labels(sign:bool|str|list, scenarios:dict) -> list[str]:
  """
  Labels `sign` asks for: a scenario of `scenarios`, `True` being `signature`, or a list as is.
  A scenario missing from `scenarios` gives one line, labelled with its name, bare for `True`.
  """
  if not sign: return []
  if isinstance(sign, (list, tuple)): return [str(x) for x in sign]
  key = "signature" if sign is True else str(sign).lower()
  return list(scenarios.get(key, [""] if sign is True else [str(sign)]))

#--------------------------------------------------------------------------------------- Signatures

def signatures(
  doc, labels:list[str], face:RichSegment, *,
  gap:float, width:float, thick:float,
  line:str = "single",
):
  """
  Signing lines after the last block, under each its label in `face`, centred.

  One line sits at the right margin; several span margin to margin, spread evenly.
  `gap` mm of blank space above them takes the handwriting, `width` mm is the widest a line gets.
  The lines are bottom borders `thick` pt thick in the colour of `face`, in a Word border style:
  `single`, or `dotted` as on a form.
  """
  if not labels: return
  def column(rule, caption, label:str, indent:float):
    """One signing column: the blank space, a bottom rule, the label under it."""
    pf = rule.paragraph_format
    pf.left_indent = Mm(indent)
    pf.space_before = Pt(mm_to_pt(gap))
    pf.space_after = Pt(0)
    _set_pbdr(rule._element, color_hex(face.color), sides=("bottom",),
      size_eighths=eighths(thick), val=line)
    caption.paragraph_format.left_indent = Mm(indent)
    caption.paragraph_format.space_before = Pt(1)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _apply_run_format(caption.add_run(label), face, face.family, face.size)
  doc._flush_para()
  content_w = doc.content_width
  count = len(labels)
  if count == 1:
    column(doc.new_para(), doc.new_para(), labels[0], max(0, content_w - width))
    return
  line_w = min(width, content_w / count - _SIGN_SPACING_MM)
  spacing = (content_w - count * line_w) / (count - 1)
  widths = [line_w if i % 2 == 0 else spacing for i in range(2 * count - 1)]
  tbl = doc.new_table(1, len(widths))
  set_table_widths(tbl, widths)
  remove_table_borders(tbl)
  cells = tbl.rows[0].cells
  for cell in cells:
    set_cell_margins(cell, top=0, right=0, bot=0, left=0)
  for i, label in enumerate(labels):
    cell = cells[2 * i]
    column(cell.paragraphs[0], cell.add_paragraph(), label, 0)

#-------------------------------------------------------------------------------------- Page number

def page_number(
  p, label:str, *,
  show_total:bool,
  face:RichSegment|None = None,
  many:bool = False,
):
  """
  `Label N/M`, or `Label N` without `show_total`, centred in footer paragraph `p`.

  The numbers are Word fields, worked out on every page.
  `face` sets their font, size and colour; without one they take the footer's style.
  With `many`, only a document of more than one page shows them.
  """
  def look(run):
    _apply_run_format(run, face, face.family, face.size)
  p.text = ""
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  template = (f"{label} " if label else "") + "{page}" + ("/{pages}" if show_total else "")
  fill = page_fields_if_many if many else page_fields
  fill(p, template, look if face else None)
